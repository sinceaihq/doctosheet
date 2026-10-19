"""CSV and XLSX export of one or more (reviewed) results.

Input is the JSON form of results (``Result.to_dict()``, possibly edited by the
user in the review interface). Amounts and quantities become numbers, dates
become dates, so spreadsheets in any locale read them correctly. Computed
suggestions are skipped unless the user accepted them. Text that a spreadsheet
could run as a formula (=, +, -, @ …) is written as plain text.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from decimal import Decimal, InvalidOperation
from typing import Any

from .model import FIELD_LABELS, HEADER_FIELDS

AMOUNT_FIELDS = {"net_amount", "vat_amount", "total_amount"}
DATE_FIELDS = {"issue_date", "due_date"}
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
LINE_COLUMNS = ("description", "quantity", "unit", "unit_price", "vat_rate", "amount")
LINE_LABELS = {"description": "Description", "quantity": "Quantity", "unit": "Unit", "unit_price": "Unit price",
               "vat_rate": "VAT %", "amount": "Amount"}  # fmt: skip
NUMERIC_LINE = {"quantity", "unit_price", "vat_rate", "amount"}


def _value(field: dict[str, Any] | None) -> str | None:
    if not field or field.get("value") in (None, ""):
        return None
    if field.get("computed"):  # suggestion not accepted by the user
        return None
    return str(field["value"])


def _typed(name: str, value: str | None) -> Any:
    if value is None:
        return None
    if name in AMOUNT_FIELDS or name in NUMERIC_LINE:
        try:
            return Decimal(value)
        except InvalidOperation:
            return value
    if name in DATE_FIELDS:
        try:
            return dt.date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _needs_review(doc: dict[str, Any]) -> str:
    low = [FIELD_LABELS[k] for k in HEADER_FIELDS if (doc["fields"].get(k) or {}).get("confidence") == "low"
           and _value(doc["fields"].get(k)) is not None]  # fmt: skip
    failed = [c["name"] for c in doc.get("checks", []) if c.get("status") == "failed"]
    return "; ".join(low + [f"check failed: {f}" for f in failed])


def rows(docs: list[dict[str, Any]]) -> tuple[list[list[Any]], list[list[Any]], list[list[Any]]]:
    """(invoice rows, line item rows, VAT rows), each with a header row first."""
    invoices = [["File", *[FIELD_LABELS[k] for k in HEADER_FIELDS], "Read from", "Needs review"]]
    lines = [["File", "Invoice number", *[LINE_LABELS[c] for c in LINE_COLUMNS]]]
    vat = [["File", "Invoice number", "VAT %", "Taxable amount", "VAT"]]
    for doc in docs:
        f = doc["fields"]
        number = _value(f.get("invoice_number"))
        invoices.append(
            [doc["file_name"], *[_typed(k, _value(f.get(k))) for k in HEADER_FIELDS], doc.get("method"),
             _needs_review(doc)]  # fmt: skip
        )
        for li in doc.get("lines", []):
            lines.append([doc["file_name"], number, *[_typed(c, _value(li.get(c))) for c in LINE_COLUMNS]])
        for row in doc.get("vat_breakdown", []):
            vat.append([doc["file_name"], number, _typed("vat_rate", _value(row.get("rate"))),
                        _typed("net_amount", _value(row.get("base"))), _typed("vat_amount", _value(row.get("vat")))])  # fmt: skip
    return invoices, lines, vat


def _safe_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def to_csv(docs: list[dict[str, Any]], table: str = "invoices", dialect: str = "standard") -> str:
    """CSV of one table. ``dialect``: "standard" (comma, 1234.56) or "eu" (semicolon, 1234,56)."""
    invoices, lines, vat = rows(docs)
    data = {"invoices": invoices, "lines": lines, "vat": vat}[table]
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";" if dialect == "eu" else ",", lineterminator="\r\n")
    for row in data:
        cells = []
        for v in row:
            if isinstance(v, Decimal):
                cells.append(f"{v}".replace(".", ",") if dialect == "eu" else f"{v}")
            elif isinstance(v, dt.date):
                cells.append(v.isoformat())
            else:
                cells.append(_safe_text(v) if v is not None else "")
        writer.writerow(cells)
    return out.getvalue()


def to_xlsx(docs: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    sheets = zip(("Invoices", "Line items", "VAT"), rows(docs))
    first = True
    for title, data in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        for r, row in enumerate(data, start=1):
            for c, v in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c)
                if isinstance(v, str):
                    cell.value = v
                    cell.data_type = "s"  # never a formula, whatever the text starts with
                else:
                    cell.value = v  # numbers stay Decimal: exact in the file
                if isinstance(v, Decimal):
                    cell.number_format = "#,##0.00" if v.as_tuple().exponent == -2 else "General"
                elif isinstance(v, dt.date):
                    cell.number_format = "yyyy-mm-dd"
        for c in range(1, len(data[0]) + 1):
            ws.cell(row=1, column=c).font = Font(bold=True)
            ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="E8ECF7")
            width = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 50)
        ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
