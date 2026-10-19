from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from doctosheet import extract
from doctosheet.einvoice import EInvoiceError, parse
from doctosheet.export import to_csv, to_xlsx
from doctosheet.extract import UnsupportedFile
from doctosheet.numbers import amounts_in, dates_in, parse_amount

SAMPLES = Path(__file__).parent / "samples"


@pytest.mark.parametrize(
    ("text", "value"),
    [
        ("1.234,56", "1234.56"),
        ("1 234,56", "1234.56"),
        ("1 234,56", "1234.56"),
        ("1,234.56", "1234.56"),
        ("1'234.56", "1234.56"),
        ("1234,56", "1234.56"),
        ("-12,50", "-12.50"),
        ("12,50-", "-12.50"),
        ("0,99", "0.99"),
    ],
)
def test_european_amounts(text: str, value: str) -> None:
    assert parse_amount(text) == Decimal(value)


@pytest.mark.parametrize("text", ["1.234.56", "1,234,56", "12.5", "1.234", "abc"])
def test_ambiguous_or_malformed_amounts_are_rejected(text: str) -> None:
    assert parse_amount(text) is None


def test_dates_are_not_amounts() -> None:
    assert [v for _, _, v in amounts_in("Zahlbar bis 04.04.2018: 235,62 EUR")] == [Decimal("235.62")]
    assert [v for _, _, v in amounts_in("VAT 25,5 % 318,75")] == [Decimal("318.75")]


@pytest.mark.parametrize(
    ("text", "date"),
    [
        ("05.03.2026", dt.date(2026, 3, 5)),
        ("5.3.2026", dt.date(2026, 3, 5)),
        ("05/03/2026", dt.date(2026, 3, 5)),
        ("2026-03-05", dt.date(2026, 3, 5)),
        ("5. maaliskuuta 2026", dt.date(2026, 3, 5)),
        ("5 mars 2026", dt.date(2026, 3, 5)),
        ("16 februari 2025", dt.date(2025, 2, 16)),
        ("4. 5. 2026", dt.date(2026, 5, 4)),
    ],
)
def test_dates_day_first(text: str, date: dt.date) -> None:
    assert [d for _, _, d in dates_in(text)] == [date]


def test_impossible_dates_are_ignored() -> None:
    assert dates_in("31.02.2026") == []


def test_text_pdf_sample() -> None:
    r = extract((SAMPLES / "invoice-fi.pdf").read_bytes(), "invoice-fi.pdf")
    v = {k: f.value for k, f in r.fields.items()}
    assert r.method == "text"
    assert v["invoice_number"] == "2026-1047"
    assert v["issue_date"] == "2026-03-15" and v["due_date"] == "2026-04-14"
    assert (v["net_amount"], v["vat_amount"], v["total_amount"]) == ("1250.00", "318.75", "1568.75")
    assert v["iban"] == "FI2112345600000785" and r.fields["iban"].confidence == "high"
    # the sample's VAT number deliberately fails its check digits: found, but flagged
    assert v["seller_vat_id"] == "FI12345672" and r.fields["seller_vat_id"].confidence == "low"
    assert [li.amount.value for li in r.lines] == ["650.00", "240.00", "360.00"]
    assert {c.name: c.status for c in r.checks}["net + VAT = total"] == "ok"
    # every value read from the page knows where it was
    assert all(f.box is not None for f in r.fields.values() if f.value is not None)


def test_embedded_einvoice_is_exact() -> None:
    data = (SAMPLES / "invoice-de-zugferd.pdf").read_bytes()
    r = extract(data, "de.pdf")
    assert r.method == "xml" and r.einvoice and "factur-x.xml" in r.einvoice
    assert r.fields["total_amount"].value == "947.08"
    assert len(r.lines) == 3 and len(r.vat_breakdown) == 2
    # the printed text gives the same totals
    printed = extract(data, "de.pdf", use_embedded_xml=False)
    assert printed.method == "text"
    assert printed.fields["total_amount"].value == "947.08"
    assert printed.fields["vat_amount"].value == r.fields["vat_amount"].value


def test_peppol_ubl() -> None:
    r = extract((SAMPLES / "invoice-nl-peppol.xml").read_bytes(), "nl.xml")
    v = {k: f.value for k, f in r.fields.items()}
    assert v["invoice_number"] == "NL-2026-552" and v["due_date"] == "2026-03-22"
    assert v["seller_name"] == "Molenaar Vertalingen B.V." and v["total_amount"] == "1452.00"
    assert r.lines[0].quantity.value == "8" and r.lines[0].unit.value == "HUR"


def test_missing_amount_is_suggested_not_filled() -> None:
    xml = (
        (SAMPLES / "invoice-nl-peppol.xml")
        .read_bytes()
        .replace(b'<cbc:TaxExclusiveAmount currencyID="EUR">1200.00</cbc:TaxExclusiveAmount>', b"")
    )
    r = extract(xml, "nl.xml")
    net = r.fields["net_amount"]
    assert net.computed and net.value == "1200.00" and net.confidence == "low"
    assert r.value("net_amount") is None  # not a real value until accepted
    # and the export leaves it out
    wb = openpyxl.load_workbook(io.BytesIO(to_xlsx([r.to_dict()])))
    header = [c.value for c in wb["Invoices"][1]]
    assert wb["Invoices"][2][header.index("Net amount")].value is None


def test_xml_attacks_are_refused() -> None:
    bomb = b"""<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;">]>
    <Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"><ID>&b;</ID></Invoice>"""
    with pytest.raises(EInvoiceError):
        parse(bomb)
    xxe = b"""<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>
    <Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"><ID>&e;</ID></Invoice>"""
    with pytest.raises(EInvoiceError):
        parse(xxe)


def test_unsupported_files() -> None:
    with pytest.raises(UnsupportedFile):
        extract(b"PK\x03\x04 not a supported file", "x.zip")


def test_export_types_and_formula_injection() -> None:
    r = extract((SAMPLES / "invoice-fi.pdf").read_bytes(), "invoice-fi.pdf").to_dict()
    r["fields"]["buyer_name"]["value"] = '=HYPERLINK("http://evil.example","x")'
    wb = openpyxl.load_workbook(io.BytesIO(to_xlsx([r])))
    ws = wb["Invoices"]
    header = [c.value for c in ws[1]]
    row = {h: c for h, c in zip(header, ws[2], strict=False)}
    assert row["Buyer"].data_type == "s"  # stored as text, never as a formula
    assert row["Total amount"].value == pytest.approx(1568.75)
    assert isinstance(row["Invoice date"].value, dt.datetime)
    csv_text = to_csv([r], "invoices", "eu")
    assert "'=HYPERLINK" in csv_text and ";=HYPERLINK" not in csv_text and ';"=HYPERLINK' not in csv_text
    assert "1568,75" in csv_text
    assert "1568.75" in to_csv([r], "invoices", "standard")


def test_receipt_photo_with_ocr() -> None:
    pytest.importorskip("rapidocr")
    r = extract((SAMPLES / "receipt-fi.png").read_bytes(), "receipt-fi.png")
    v = {k: f.value for k, f in r.fields.items()}
    assert r.method == "ocr"
    assert v["invoice_number"] == "88123" and v["issue_date"] == "2026-03-12"
    assert v["total_amount"] == "17.90" and v["vat_amount"] == "2.20"
    # the net amount is not printed on the receipt: suggested, not filled in
    assert r.fields["net_amount"].computed and r.value("net_amount") is None
    assert [li.amount.value for li in r.lines] == ["8.40", "3.50", "6.00"]
